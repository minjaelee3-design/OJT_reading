import datetime
import io
import json
import os
import re
import time
from typing import Dict, List, Optional, Tuple

import streamlit as st
from dotenv import load_dotenv
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# .env 환경 변수 자동 로드
load_dotenv()

# ==========================================
# 1. 페이지 환경설정 및 커스텀 스타일링
# ==========================================
st.set_page_config(
    page_title="신입사원 독서경영 IDEA SHEET 자동 생성기",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    .main-title {
        font-size: 2.1rem;
        font-weight: 700;
        color: #1E3A8A;
        margin-bottom: 0.2rem;
    }
    .sub-title {
        font-size: 1.05rem;
        color: #4B5563;
        margin-bottom: 1.5rem;
    }
    .card-box {
        background-color: #F8FAFC;
        border: 1px solid #E2E8F0;
        border-radius: 10px;
        padding: 1.2rem;
        margin-bottom: 1rem;
    }
    .badge-success {
        display: inline-block;
        background-color: #DCFCE7;
        color: #166534;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 600;
        font-size: 0.88rem;
    }
    .badge-warning {
        display: inline-block;
        background-color: #FEF3C7;
        color: #92400E;
        padding: 4px 10px;
        border-radius: 6px;
        font-weight: 600;
        font-size: 0.88rem;
    }
    .stButton>button {
        border-radius: 8px;
        font-weight: 600;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ==========================================
# 2. Gemini API 호출 유틸리티
# ==========================================
def call_gemini_api(api_key: str, prompt: str, model_name: str = "gemini-3.8-flash") -> str:
    """
    최신 google-genai SDK를 호출합니다.
    일일 한도(20회) 및 분당 한도(5회)를 구분하여 정확한 한글 안내를 반환합니다.
    """
    if not api_key:
        raise ValueError("Google Gemini API 키가 입력되지 않았습니다. 사이드바에서 API 키를 입력해 주세요.")

    try:
        from google import genai
    except ImportError:
        raise RuntimeError("google-genai 패키지가 설치되어 있지 않습니다. requirements.txt 설치를 확인해 주세요.")

    client = genai.Client(api_key=api_key)

    try:
        response = client.models.generate_content(
            model=model_name,
            contents=prompt,
        )
        if response and response.text:
            return response.text
        else:
            raise ValueError("Gemini API로부터 빈 텍스트 응답이 반환되었습니다.")

    except Exception as e:
        err_msg = str(e)

        # 1. 일일 사용 한도 (20회/일) 초과 여부 확인
        if "PerDay" in err_msg or "limit: 20" in err_msg:
            raise RuntimeError("오늘 무료 일일 사용량(20회)을 모두 소진했습니다. AI Studio에서 '새 프로젝트'를 만들어 새 키를 발급받으세요.") from e

        # 2. 분당 호출 한도 (5회/분) 초과 확인
        if any(code in err_msg for code in ["429", "ResourceExhausted", "RESOURCE_EXHAUSTED"]):
            raise RuntimeError("분당 호출 한도(5회)에 일시 도달했습니다. 1분 뒤 다시 시도해 주세요.") from e

        # 3. 구글 서버 일시 지연 (503)
        if any(code in err_msg for code in ["503", "UNAVAILABLE", "Overloaded"]):
            raise RuntimeError("구글 AI 서버 접속량이 일시적으로 많습니다. 30초 뒤 다시 시도해 주세요.") from e

        raise e


@st.cache_data(show_spinner=False, ttl=3600)
def fetch_book_info_and_summary(api_key: str, book_title: str, model_name: str) -> Dict:
    """
    도서명을 기반으로 요약을 추출하며, 동일 도서는 1시간 동안 API 호출 없이 캐시를 재사용합니다.
    """
    prompt = f"""
당신은 기업 사내 독서경영 및 도서 분석 전문가입니다.
사용자가 입력한 도서명 '{book_title}'에 대한 정보를 바탕으로 다음 항목을 JSON 형식으로 작성해 주세요.

반드시 마크다운 코드블록(```json ... ```) 또는 순수 JSON 형식으로 응답하세요.

JSON 필드 규격:
{{
  "author": "저자명 (외국 저자인 경우 한국어 표기 및 원어 병기 가능)",
  "category": "분야/카테고리 (예: 자기계발, 경영전략, 생산/품질, 기술 트렌드 등)",
  "overview": "도서 전체의 핵심 주제 및 전달 가치 요약 (3~4문장)",
  "chapters": [
    {{
      "theme": "핵심 챕터 또는 주제 1 (예: 1장 또는 핵심 테마명)",
      "summary": "해당 챕터/테마의 주요 핵심 내용 및 교훈 요약 (2~3문장)"
    }},
    {{
      "theme": "핵심 챕터 또는 주제 2",
      "summary": "해당 챕터/테마의 주요 핵심 내용 및 교훈 요약 (2~3문장)"
    }},
    {{
      "theme": "핵심 챕터 또는 주제 3",
      "summary": "해당 챕터/테마의 주요 핵심 내용 및 교훈 요약 (2~3문장)"
    }},
    {{
      "theme": "핵심 챕터 또는 주제 4",
      "summary": "해당 챕터/테마의 주요 핵심 내용 및 교훈 요약 (2~3문장)"
    }}
  ]
}}
"""
    raw_response = call_gemini_api(api_key, prompt, model_name)

    cleaned = raw_response.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)

    try:
        data = json.loads(cleaned)
        return data
    except Exception:
        author_match = re.search(r'"author"\s*:\s*"([^"]+)"', raw_response)
        author = author_match.group(1) if author_match else "저자 미상"
        return {
            "author": author,
            "category": "일반",
            "overview": raw_response[:300] + "...",
            "chapters": [
                {"theme": "도서 핵심 내용 요약", "summary": raw_response}
            ],
        }


def generate_idea_sheet_content(
    api_key: str,
    book_title: str,
    author: str,
    department: str,
    user_name: str,
    book_summary: str,
    user_notes: str,
    user_ideas: str,
    model_name: str,
) -> str:
    """
    단 1회 호출로 15줄 이상의 비즈니스 독서 보고서를 완벽히 작성합니다 (쿼터 낭비 제거).
    """
    has_summary = bool(book_summary and book_summary.strip())
    has_author = bool(author and author.strip())
    has_user_ideas = bool(user_ideas and user_ideas.strip())
    has_user_notes = bool(user_notes and user_notes.strip())

    if has_summary:
        book_info_prompt = f"""[도서 정보 및 기확보 요약]
- 도서명: {book_title}
- 저자명: {author if has_author else '저자 미상'}
- 도서 핵심 요약:
{book_summary.strip()}"""
    else:
        book_info_prompt = f"""[도서 정보 (사전 요약 미제공)]
- 도서명: {book_title}
- 저자명: {author if has_author else 'AI 지식으로 자동 유추'}

★ [AI 도서 유추 지침]: 사용자가 사전에 도서 검색/요약을 거치지 않았으므로, Gemini가 보유한 도서 지식을 총동원하여 '{book_title}'(저자: {author if has_author else '해당 도서의 저자'})의 핵심 주제, 출판 배경, 핵심 논지 및 챕터별 주요 시사점을 스스로 파악·유추하여 풍부하게 작성하세요."""

    instructions = []
    if has_user_ideas:
        instructions.append(f"""- [사용자 개선 아이디어 반영]: 사용자가 직접 제시한 아래 아이디어를 보고서의 2단계(업무 개선안)에 핵심으로 채택하여 현업 적용 방안을 구체화하세요:\n'''{user_ideas.strip()}'''""")
    else:
        instructions.append(f"""- [신입사원 관점 업무 개선 아이디어 독자 도출]: 사용자의 별도 요청 아이디어가 공란이므로, {department} 신입사원({user_name})의 시각에서 도서의 교훈과 핵심 방법론을 현업에 접목할 수 있는 실질적인 업무 개선 아이디어를 최소 2가지 이상 구체적으로 창안하여 제시하세요.
  (예: 현장 공정 낭비 요소 제거, 품질 관리 프로세스 표준화, 5S 및 안전 관리, 디지털/스마트 툴을 활용한 인수인계 효율화, 팀 내 소통 활성화 등 {department}에 실질적 도움이 되는 구체적 방안)""")

    if has_user_notes:
        instructions.append(f"""- [사용자 독서 메모 인용]: 사용자가 독서 중 기록한 아래 메모/인상 깊은 구절을 내용정리의 적절한 위치에 자연스럽게 인용하거나 반영하세요:\n'''{user_notes.strip()}'''""")

    extra_instructions = "\n".join(instructions)

    prompt = f"""당신은 대기업 사내 독서경영 코치이자 비즈니스 문서 작성 전문가입니다.
사내 공식 양식인 [독서경영 IDEA SHEET]의 '7. 내용정리' 란에 기입할 최고 품질의 비즈니스 독서 보고서를 작성해야 합니다.

[작성자 및 부서 정보]
- 부서명: {department}
- 작성자(성명): {user_name}

{book_info_prompt}

{extra_instructions}

[반드시 준수해야 하는 엄격한 제약 사항]
1. 분량 기준 (최우선 필수):
   - 결과 텍스트는 엔터(개행 문자 \\n) 기준 **반드시 16줄 이상 22줄 이하의 풍성한 줄글 또는 문단 구조**로 작성해야 합니다.
   - 각 소주제(1. 도서 핵심 요약, 2. 업무 개선 아이디어, 3. 실천 다짐 및 느낀 점)별로 구체적인 설명과 개행을 풍성하게 두어 줄 수를 충분히 확보하세요.
2. 내용 구성 (3단계 프레임워크 필수 포함):
   - [1단계: 도서 핵심 요약 및 통찰]: 책이 제시하는 핵심 메시지와 주요 원리 (최소 5줄 이상)
   - [2단계: 실무 프로세스 개선 아이디어 제안]: {department} 신입사원 시각에서 현업에 실제 적용할 수 있는 구체적 개선안 및 기대효과 (최소 6줄 이상)
   - [3단계: 신입사원 실천 다짐 및 느낀 점]: 현업 수행 자세, 팀워크, 개인적 성장 목표에 대한 다짐 (최소 4줄 이상)
3. 어조 및 문체:
   - 비즈니스에 적합한 단정하고 논리적이며 성실함이 묻어나는 보고서 문체 (~하였습니다, ~하고자 합니다, ~하겠습니다).
4. 출력 형식:
   - 불필요한 서두/결미 멘트는 제외하고 오직 '7. 내용정리' 칸에 그대로 들어갈 본문 텍스트만 출력하세요.
"""
    result = call_gemini_api(api_key, prompt, model_name)
    cleaned_result = result.strip() if result else ""

    # 추가 API 재호출(쿼터 낭비)을 완전히 없애고 첫 결과 바로 반환
    return cleaned_result


# ==========================================
# 3. PPTX 템플릿 처리 모듈
# ==========================================
def find_target_table(prs: Presentation):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                full_text = "".join(cell.text for row in table.rows for cell in row.cells)
                if "부서명" in full_text or "내용정리" in full_text:
                    return slide, table
    return None, None


def fill_cell_text(cell, text: str, font_size_pt: int = 11, bold: bool = False, align=PP_ALIGN.CENTER, font_name: str = "맑은 고딕"):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text) if text else ""
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    run.font.name = font_name
    run.font.color.rgb = RGBColor(0, 0, 0)


def fill_content_cell(cell, content_text: str, font_size_pt: int = 9, font_name: str = "맑은 고딕"):
    cell.text = ""
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.08)
    cell.margin_bottom = Inches(0.08)

    lines = content_text.split("\n")
    for idx, line in enumerate(lines):
        if idx == 0:
            p = cell.text_frame.paragraphs[0]
        else:
            p = cell.text_frame.add_paragraph()

        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.15
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size_pt)
        run.font.name = font_name
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def generate_completed_pptx(template_source, department: str, user_name: str, book_title: str, author: str, start_date: str, end_date: str, content_text: str) -> io.BytesIO:
    prs = Presentation(template_source)
    _, table = find_target_table(prs)

    if table is None:
        raise ValueError("PPTX 파일 내에서 독서경영 양식 표를 찾을 수 없습니다. 올바른 템플릿인지 확인해 주세요.")

    fill_cell_text(table.cell(0, 1), department, font_size_pt=11, bold=False, align=PP_ALIGN.CENTER)
    fill_cell_text(table.cell(0, 3), user_name, font_size_pt=11, bold=False, align=PP_ALIGN.CENTER)
    fill_cell_text(table.cell(1, 1), book_title, font_size_pt=11, bold=False, align=PP_ALIGN.CENTER)
    fill_cell_text(table.cell(1, 3), author, font_size_pt=11, bold=False, align=PP_ALIGN.CENTER)
    fill_cell_text(table.cell(2, 1), start_date, font_size_pt=11, bold=False, align=PP_ALIGN.CENTER)
    fill_cell_text(table.cell(2, 3), end_date, font_size_pt=11, bold=False, align=PP_ALIGN.CENTER)
    fill_content_cell(table.cell(4, 0), content_text, font_size_pt=9)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output


# ==========================================
# 4. Streamlit 세션 상태 초기화
# ==========================================
DEFAULT_TEMPLATE_PATH = "신입사원_독서경영_양식 (0주차_소속부서_도서명).pptx"
if not os.path.exists(DEFAULT_TEMPLATE_PATH):
    DEFAULT_TEMPLATE_PATH = "신입사원_독서경영_양식 (1주차_대경생산팀_도서명).pptx"

if "book_info" not in st.session_state:
    st.session_state.book_info = None
if "content_summary" not in st.session_state:
    st.session_state.content_summary = ""
if "book_title_input" not in st.session_state:
    st.session_state.book_title_input = ""
if "author_input" not in st.session_state:
    st.session_state.author_input = ""
if "start_date_str" not in st.session_state:
    st.session_state.start_date_str = "2026.10.13"
if "end_date_str" not in st.session_state:
    st.session_state.end_date_str = "2026.10.16"
if "final_content_editor" not in st.session_state:
    st.session_state.final_content_editor = ""


# ==========================================
# 5. 사이드바 구성
# ==========================================
with st.sidebar:
    st.title("⚙️ 설정 및 환경")

    env_api_key = os.getenv("GEMINI_API_KEY", "")
    api_key = st.text_input(
        "Google Gemini API Key",
        value=env_api_key,
        type="password",
        help="Google AI Studio([https://aistudio.google.com/)에서](https://aistudio.google.com/)에서) 발급받은 API 키를 입력하세요.",
    )
    if not api_key:
        st.warning("⚠️ Gemini API 키를 입력하거나 .env 파일에 등록해 주세요.")
    else:
        st.success("🔑 API 키 준비 완료")

    model_choice = st.selectbox(
        "사용 AI 모델",
        ["gemini-3.8-flash"],
        index=0,
    )

    st.markdown("---")
    st.subheader("📄 PPTX 템플릿 관리")

    has_default_template = os.path.exists(DEFAULT_TEMPLATE_PATH)
    if has_default_template:
        st.info(f"✅ 기본 템플릿 로드됨:\n`{DEFAULT_TEMPLATE_PATH}`")
    else:
        st.warning("기본 템플릿 파일을 찾을 수 없습니다. 아래에서 업로드해 주세요.")

    uploaded_template = st.file_uploader(
        "사용자 정의 PPTX 템플릿 업로드 (선택)",
        type=["pptx"],
    )


# ==========================================
# 6. 메인 화면 UI
# ==========================================
st.markdown('<div class="main-title">📚 신입사원 독서경영 IDEA SHEET 자동 생성기</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="sub-title">밀리의 서재 등 독서 후, Gemini AI로 챕터 요약과 업무 개선 아이디어를 도출하여 사내 공식 PPTX 서식에 즉시 주입합니다.</div>',
    unsafe_allow_html=True,
)

tab1, tab2 = st.tabs(["🚀 IDEA SHEET 생성 워크스페이스", "📋 양식 및 규칙 가이드"])

with tab1:
    st.markdown("### 1️⃣ 도서 검색 및 핵심 요약 추출 (선택)")
    col_search1, col_search2 = st.columns([4, 1])

    with col_search1:
        searched_book = st.text_input(
            "도서명 입력",
            value=st.session_state.book_title_input,
            placeholder="예: 원칙(Principles), 피터 드러커의 최고의 질문, 도커 교과서 등",
            key="book_search_field",
        )
    with col_search2:
        st.write("")
        st.write("")
        search_clicked = st.button("🔍 도서 정보 조회", use_container_width=True)

    if search_clicked:
        if not searched_book.strip():
            st.error("도서명을 먼저 입력해 주세요.")
        elif not api_key:
            st.error("Gemini API 키가 필요합니다. 사이드바에 API 키를 입력해 주세요.")
        else:
            with st.spinner(f"'{searched_book}'의 저자 정보 및 챕터별 핵심 내용을 탐색 중입니다..."):
                try:
                    info = fetch_book_info_and_summary(api_key, searched_book, model_choice)
                    st.session_state.book_info = info
                    st.session_state.book_title_input = searched_book
                    st.session_state.author_input = info.get("author", "")
                    st.success(f"도서 정보 탐색 완료! (저자: {st.session_state.author_input})")
                except Exception as e:
                    st.error(f"❌ 오류: {str(e)}")

    if st.session_state.book_info:
        info = st.session_state.book_info
        with st.expander("📖 AI 분석 도서 개요 및 챕터별 핵심 내용 확인", expanded=True):
            st.markdown(f"**카테고리**: `{info.get('category', '일반')}` | **저자**: `{info.get('author', '저자 미상')}`")
            st.info(f"**전체 개요**: {info.get('overview', '')}")

            chapters = info.get("chapters", [])
            if chapters:
                st.markdown("**주요 챕터 및 테마 요약:**")
                ch_cols = st.columns(min(len(chapters), 3) if len(chapters) <= 3 else 2)
                for idx, ch in enumerate(chapters):
                    col_target = ch_cols[idx % len(ch_cols)]
                    with col_target:
                        st.markdown(
                            f"""
                            <div class="card-box">
                                <b style="color: #1E40AF;">📌 {ch.get('theme', f'주제 {idx+1}')}</b><br/>
                                <span style="font-size: 0.92rem; color: #374151;">{ch.get('summary', '')}</span>
                            </div>
                            """,
                            unsafe_allow_html=True,
                        )

    user_notes = st.text_area(
        "📝 독서 중 메모 또는 인상 깊은 문장 (선택사항)",
        placeholder="밀리의 서재를 읽으며 기록한 주요 발췌문, 감명받은 구절을 적어주세요.",
        height=100,
    )

    st.markdown("---")

    st.markdown("### 2️⃣ PPT 양식 기입 정보 설정")

    col_meta1, col_meta2, col_meta3 = st.columns(3)
    with col_meta1:
        week_value = st.selectbox("주차 (파일명 생성용)", ["1주차", "2주차", "3주차", "4주차", "5주차", "6주차", "7주차", "8주차"], index=0)
    with col_meta2:
        department_value = st.text_input("부서명", value="대경생산팀")
    with col_meta3:
        user_name_value = st.text_input("성명", value="이민재")

    col_book1, col_book2 = st.columns(2)
    with col_book1:
        final_book_title = st.text_input("도서명", value=st.session_state.book_title_input)
    with col_book2:
        final_author = st.text_input("저자명", value=st.session_state.author_input)

    col_date1, col_date2 = st.columns(2)
    with col_date1:
        st.markdown("**5. 시작일**")
        d_col1, d_col2 = st.columns([1, 1])
        with d_col1:
            try:
                parts = [int(p) for p in st.session_state.start_date_str.split(".")]
                default_start = datetime.date(parts[0], parts[1], parts[2])
            except Exception:
                default_start = datetime.date.today()
            picker_start = st.date_input("시작일 달력 선택", value=default_start, key="picker_start")
        with d_col2:
            start_date_input = st.text_input(
                "시작일 텍스트 (YYYY.MM.DD)",
                value=picker_start.strftime("%Y.%m.%d"),
                key="text_start_date",
            )

    with col_date2:
        st.markdown("**6. 완독일**")
        d_col3, d_col4 = st.columns([1, 1])
        with d_col3:
            try:
                parts = [int(p) for p in st.session_state.end_date_str.split(".")]
                default_end = datetime.date(parts[0], parts[1], parts[2])
            except Exception:
                default_end = datetime.date.today()
            picker_end = st.date_input("완독일 달력 선택", value=default_end, key="picker_end")
        with d_col4:
            end_date_input = st.text_input(
                "완독일 텍스트 (YYYY.MM.DD)",
                value=picker_end.strftime("%Y.%m.%d"),
                key="text_end_date",
            )

    user_ideas_prompt = st.text_area(
        "💡 업무 개선 아이디어 및 특별 요청사항 (선택사항)",
        placeholder="공란으로 둘 경우, 도서 내용을 바탕으로 신입사원 시각의 실무 개선 방안이 자동 구성됩니다.",
        height=90,
    )

    st.markdown("---")

    st.markdown("### 3️⃣ '7. 내용정리' 생성 및 편집 (15줄 이상, 9pt)")

    gen_col1, gen_col2 = st.columns([2, 5])
    with gen_col1:
        create_content_btn = st.button("✨ 7. 내용정리 AI 자동 생성", type="primary", use_container_width=True)

    if create_content_btn:
        effective_book_title = final_book_title.strip() or searched_book.strip()
        effective_author = final_author.strip()

        if not effective_book_title:
            st.error("❌ 도서명이 입력되지 않았습니다. 2단계 도서명을 입력해 주세요.")
        elif not api_key:
            st.error("❌ Gemini API 키가 필요합니다. 사이드바에 API 키를 입력해 주세요.")
        else:
            with st.spinner("15줄 이상의 비즈니스 독서경영 보고서를 작성 중입니다..."):
                try:
                    summary_context = ""
                    if st.session_state.book_info:
                        info_dict = st.session_state.book_info
                        overview = info_dict.get("overview", "")
                        chapters = info_dict.get("chapters", [])
                        if isinstance(chapters, list):
                            ch_texts = [
                                f"[{c.get('theme', '주제')}]: {c.get('summary', '')}"
                                for c in chapters
                                if isinstance(c, dict)
                            ]
                            summary_context = f"{overview}\n" + "\n".join(ch_texts)
                        else:
                            summary_context = str(overview)

                    generated_text = generate_idea_sheet_content(
                        api_key=api_key,
                        book_title=effective_book_title,
                        author=effective_author,
                        department=department_value,
                        user_name=user_name_value,
                        book_summary=summary_context,
                        user_notes=user_notes,
                        user_ideas=user_ideas_prompt,
                        model_name=model_choice,
                    )

                    if not generated_text or not generated_text.strip():
                        st.error("❌ Gemini API로부터 응답이 비어 있습니다.")
                    else:
                        st.session_state.content_summary = generated_text
                        st.session_state["final_content_editor"] = generated_text
                        st.toast("내용정리 작성이 완료되었습니다!", icon="🎉")
                        st.rerun()
                except Exception as e:
                    st.error(f"❌ {str(e)}")

    edited_content = st.text_area(
        "최종 기입될 내용정리 본문 (수정 가능)",
        value=st.session_state.content_summary,
        height=360,
        help="생성된 문구를 직접 보완하거나 다듬을 수 있습니다.",
        key="final_content_editor",
    )
    if edited_content != st.session_state.content_summary:
        st.session_state.content_summary = edited_content

    non_empty_lines = [l for l in edited_content.split("\n") if l.strip()]
    line_count = len(non_empty_lines)
    char_count = len(edited_content)

    col_stat1, col_stat2, col_stat3 = st.columns([2, 2, 4])
    with col_stat1:
        if line_count >= 15:
            st.markdown(f'<div class="badge-success">✅ {line_count}줄 작성됨 (15줄 이상 충족)</div>', unsafe_allow_html=True)
        elif line_count > 0:
            st.markdown(f'<div class="badge-warning">⚠️ 현재 {line_count}줄 (15줄 이상 필요)</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="badge-warning">내용이 비어 있습니다.</div>', unsafe_allow_html=True)

    with col_stat2:
        st.caption(f"총 {char_count:,}자")

    st.markdown("---")

    st.markdown("### 4️⃣ 완성된 PPTX 다운로드")

    actual_book_title = final_book_title.strip() or searched_book.strip()
    clean_book_name = re.sub(r'[\\/*?:"<>|]', "", actual_book_title) or "도서명"
    output_filename = f"신입사원_독서경영_양식 ({week_value}_{department_value}_{clean_book_name}).pptx"

    template_file_to_use = uploaded_template if uploaded_template is not None else (DEFAULT_TEMPLATE_PATH if os.path.exists(DEFAULT_TEMPLATE_PATH) else None)

    col_down1, col_down2 = st.columns([2, 4])
    with col_down1:
        if template_file_to_use is None:
            st.error("PPTX 템플릿 파일이 없습니다. 사이드바에서 업로드해 주세요.")
        elif not edited_content.strip():
            st.info("내용정리를 먼저 생성하거나 작성해 주세요.")
        else:
            try:
                pptx_bytes = generate_completed_pptx(
                    template_source=template_file_to_use,
                    department=department_value,
                    user_name=user_name_value,
                    book_title=final_book_title,
                    author=final_author,
                    start_date=start_date_input,
                    end_date=end_date_input,
                    content_text=edited_content,
                )

                st.download_button(
                    label="📥 완성된 PPTX 다운로드",
                    data=pptx_bytes,
                    file_name=output_filename,
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    type="primary",
                    use_container_width=True,
                )
            except Exception as e:
                st.error(f"PPTX 생성 중 오류: {str(e)}")

    with col_down2:
        st.caption(f"📌 저장 파일명: `{output_filename}`")

with tab2:
    st.markdown("### 📋 독서경영 IDEA SHEET 규격 안내")
    st.markdown("- 15줄 이상, 9pt 폰트로 서식에 정확히 자동 반영됩니다.")
